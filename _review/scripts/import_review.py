"""Import the manually reviewed workbook into _data/tools_v2.yml.

Colour codes (from the reviewers): green = verified, yellow = verified after a fix, white = not
checked yet, purple (and pink) = still to decide, strikethrough = must not be added.
Verified values go to the top level of an entry; unchecked and undecided values go to `pending:`
(hidden on the live site, shown in preview). Struck values and struck tools are dropped.
"""
import datetime, json, re, sys
from pathlib import Path
import yaml
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[2]
XLSX = REPO / "_review/tools4fair-annotation-form_new_v2.xlsx"
OUT = REPO / "_data/tools_v2.yml"

wb = load_workbook(XLSX)
tw = wb["Tools"]
# Fields that were not part of the review (institutes, FAIRsharing id) are carried over from the last
# version before the first import, not from OUT: re-running the import must not lose them.
BEFORE = Path(__file__).with_name("tools_v2_before_import.yml")
old = {e["id"]: e for e in yaml.safe_load(open(BEFORE))}
# Fields the sheet does not (or no longer) hold come from the current file: old page URLs
# (redirect_from) and related tools (the Related columns were removed from the sheet)
current = {e["id"]: e for e in yaml.safe_load(open(OUT))}

def colour(c):
    f = c.fill
    if not f or not f.fill_type:
        return "white"
    if f.fgColor.type == "rgb":
        return {"FF92D050": "green", "FFFFFF00": "yellow", "FF7030A0": "purple", "FFFF5DBB": "purple",
                "FFF8D0D0": "broken", "FFEEF0F3": "white"}.get(f.fgColor.rgb, "white")
    return "purple" if f.fgColor.theme not in (0, None) else "white"   # theme 0 = plain white fill

def status(c):
    if c.value in (None, "") or (isinstance(c.value, str) and c.value.strip() in ("", "None")):
        return None
    if c.font and c.font.strike:
        return "struck"
    k = colour(c)
    return {"green": "verified", "yellow": "verified", "purple": "decide", "broken": "pending"}.get(k, "pending")

COL = {}
for j in range(1, tw.max_column + 1):
    COL.setdefault(str(tw.cell(row=2, column=j).value).strip(), j)
def cols(first, last):
    a, b = COL[first], COL[last]
    return list(range(a, b + 1))
# multi-value fields include the reviewers' extra unlabelled columns
TYPE_COLS = list(range(COL["Tool type 1"], COL["F"]))           # F..L
DOMAIN_COLS = cols("Domain 1", "Domain 5")
LICENCE_COLS = cols("Licence 1", "Licence 2")
CODE_COLS = cols("Code availability 1", "Code availability 2")
FAIR = [("F", "F"), ("A", "A"), ("I", "I"), ("R", "R"), ("Assessment & planning", "planning")]
STAGES = ["Plan", "Collect", "Process", "Analyse", "Preserve", "Share", "Reuse"]

# ---- reviewers' notes in Documentation / Download, turned into structured values ----------------
DOCS_NOTES = {   # text in a verified Documentation cell -> (docs URL or None, docs_note)
    "platform website": (None, "On the website"),
    "built in": (None, "Built into the tool"),
    "website and github documentation": (None, "On the website and in the GitHub repository"),
    "platform-documentatie download underneath left in page bottom": (None, "Download it at the bottom of the product page"),
    "github readme": ("{homepage}#readme", None),
}
SAAS = re.compile(r"^saas\b", re.I)
DOWNLOAD_SPECIAL = {   # tool id -> download links for verified cells that hold several labelled URLs
    "topbraid-shacl-api": [{"label": "GitHub releases", "url": "https://github.com/TopQuadrant/shacl/releases"}],
}
TOOL_IDS = {"Sextans-Suite": "sextans-suite"}
# editors' changes after the review, kept on re-import; listed on the status page so the sheet can follow
LEFT_OUT = {"sextans-suite"}                                              # not in the catalogue for now
CHANGES = {   # tool id -> {field: verified value instead of the sheet's}; "links.<name>" for a link field
    "yasgui": {"logo": "https://avatars.githubusercontent.com/u/7573184?v=4"},
    "topbraid-shacl-api": {"logo": "https://media.licdn.com/dms/image/v2/C4E0BAQGZbqII1tkEEQ/company-logo_200_200/company-logo_200_200/0/1662058039652/topquadrant_logo?e=2147483647&v=beta&t=KLfeEeC2EBFwWu7O5XPB5KSR3o7mxcsE5Zt-JBrDtO4"},
    # OpenRefine runs on your own computer (in your browser) and is not offered online: openrefine.org is the
    # project website, and the web, plug-in and workbench types misled; SPARQL only comes with extensions
    "openrefine": {"links.homepage_opens_tool": False,
                   "tool_type": ["Desktop application"],
                   "short_description": "Desktop tool, used in your browser, to clean messy tables, match values to external identifiers such as Wikidata, and convert tables to RDF with an extension.",
                   "keywords": ["CSV", "RDF"]},
    # the sheet's download link went to the WebProtégé wiki; desktop Protégé is downloaded from the software page
    "protege": {"links.download": "https://protege.stanford.edu/software/#desktop-protege",
                "links.other": [{"kind": "Report an issue", "url": "https://github.com/protegeproject/protege/issues"}]},
}
KEYWORDS = {"digital object identifier": "DOI", "doi": "DOI"}   # sheet spelling -> _data/keywords.yml
NO_VALUE = {"unknown"}                                            # a reviewed cell saying there is no value

report = {"removed_tools": [], "left_out": [], "overrides": [], "struck_values": [], "decide": [], "notes_mapped": []}
entries = []
for r in range(3, tw.max_row + 1):
    raw_id = tw.cell(row=r, column=COL["Tool id"]).value
    if not raw_id:
        continue
    if tw.cell(row=r, column=COL["Tool id"]).font.strike:
        report["removed_tools"].append(raw_id); continue
    tid = TOOL_IDS.get(raw_id, raw_id)
    if tid in LEFT_OUT:
        report["left_out"].append(raw_id); continue
    prev = old.get(tid, {})
    top, pend = {"id": tid}, {}
    if current.get(tid, {}).get("redirect_from"): top["redirect_from"] = current[tid]["redirect_from"]
    def cell(name): return tw.cell(row=r, column=COL[name])
    def put(key, value, st, note_field=None):
        if st == "verified":
            top[key] = value
        elif st in ("pending", "decide"):
            pend[key] = value
            if st == "decide": report["decide"].append((tid, note_field or key, value))
    def scalar(name, key, conv=lambda v: str(v).strip()):
        c = cell(name); st = status(c)
        if st == "struck": report["struck_values"].append((tid, name, c.value)); return
        if st: put(key, conv(c.value), st, name)
    def multi(colnums, key, conv=lambda v: str(v).strip()):
        for j in colnums:
            c = tw.cell(row=r, column=j); st = status(c)
            if st == "struck": report["struck_values"].append((tid, key, c.value)); continue
            if not st or str(c.value).strip().lower() in NO_VALUE: continue
            v = conv(c.value)
            target = top if st == "verified" else pend
            if v not in target.setdefault(key, []): target[key].append(v)
            if st == "decide": report["decide"].append((tid, key, v))

    scalar("Name (tool name)", "name")
    scalar("Short description", "short_description")
    scalar("About text", "description")
    multi(TYPE_COLS, "tool_type")
    # FAIR tags and life cycle are Yes columns
    for head, tag in FAIR:
        c = cell(head); st = status(c)
        if st == "struck": report["struck_values"].append((tid, head, c.value)); continue
        if st: (top if st == "verified" else pend).setdefault("fair", []).append(tag)
    for s in STAGES:
        c = cell(s); st = status(c)
        if st == "struck": report["struck_values"].append((tid, s, c.value)); continue
        if st: (top if st == "verified" else pend).setdefault("lifecycle", []).append(s)
    multi(DOMAIN_COLS, "domains")
    # a licence cell can hold a reviewer note such as "empty is now unconfirmed"
    multi(LICENCE_COLS, "license", lambda v: "unconfirmed" if "unconfirmed" in str(v).lower() else re.sub(r",\s*and:?$", "", str(v).strip()))
    scalar("Cost", "cost", lambda v: {"unconfirmed": "Unconfirmed"}.get(str(v).strip().lower(), str(v).strip()))
    multi(CODE_COLS, "code_availability")
    scalar("Maturity", "maturity")
    scalar("Legacy/deprecated", "deprecated", lambda v: True)

    # links
    links_top, links_pend = {}, {}
    c = cell("Homepage"); st = status(c)
    homepage = str(c.value).strip() if st else None
    if st == "verified": links_top["homepage"] = homepage
    elif st in ("pending", "decide"):
        links_pend["homepage"] = homepage; report["decide"].append((tid, "Homepage", homepage))
    c = cell("homepage link: komt direct uit bij werkende tool/software?"); st = status(c)
    if st:
        yes = str(c.value).strip().lower() in ("yes", "ja", "y", "j")
        (links_top if st == "verified" else links_pend)["homepage_opens_tool"] = yes
        if st == "decide": report["decide"].append((tid, "Homepage opens the tool", c.value))
    # documentation
    c = cell("Documentation"); st = status(c)
    if st:
        v = str(c.value).strip()
        m = re.fullmatch(r"(https?://\S+)\s*(\(.*\))?", v)   # a URL, possibly with a remark such as "(currently down)"
        if m:
            (links_top if st == "verified" else links_pend)["docs"] = m.group(1)
            if m.group(2): report["notes_mapped"].append((tid, "Documentation", v, m.group(1)))
        else:
            url, note = DOCS_NOTES[v.lower()]
            target = links_top if st == "verified" else links_pend
            if url: target["docs"] = url.format(homepage=homepage.rstrip("/"))
            if note: target["docs_note"] = note
            report["notes_mapped"].append((tid, "Documentation", v, url.format(homepage=homepage.rstrip("/")) if url else note))
        if st == "decide": report["decide"].append((tid, "Documentation", v))
    # download / install / distribution
    c = cell("Download / Install / Distribution"); st = status(c)
    if st:
        v = str(c.value).strip()
        target = links_top if st == "verified" else links_pend
        if st == "decide":
            target["download_note"] = v; report["decide"].append((tid, "Download", v))
        elif tid in DOWNLOAD_SPECIAL:
            if SAAS.match(v): top["hosted"] = True
            target["download"] = DOWNLOAD_SPECIAL[tid]
            report["notes_mapped"].append((tid, "Download", v, DOWNLOAD_SPECIAL[tid]))
        elif SAAS.match(v):
            top["hosted"] = True
            report["notes_mapped"].append((tid, "Download", v, "hosted: true"))
        elif v.lower().startswith("nvt"):
            report["notes_mapped"].append((tid, "Download", v, "not applicable: no download"))
        elif re.fullmatch(r"https?://\S+", v):
            target["download"] = v
        else:
            sys.exit(f"unmapped download note for {tid}: {v!r}")
    c = cell("Logo URL"); st = status(c)
    if st == "struck": report["struck_values"].append((tid, "Logo", c.value))
    elif st == "verified": top["logo"] = str(c.value).strip()
    elif st:
        pend["logo"] = str(c.value).strip()
        if st == "decide": report["decide"].append((tid, "Logo", c.value))
    # other links: "Kind — URL", or a bare URL
    for j in cols("Other link 1", "Other link 7"):
        c = tw.cell(row=r, column=j); st = status(c)
        if not st or st == "struck":
            if st == "struck": report["struck_values"].append((tid, "Other link", c.value))
            continue
        v = str(c.value).strip()
        kind, url = (v.split(" — ", 1) if " — " in v else ("Link", v))
        url = url.strip().rstrip(",;")                      # "https://…, " in a cell
        (links_top if st == "verified" else links_pend).setdefault("other", []).append({"kind": kind.strip(), "url": url})
    if links_top: top["links"] = links_top
    if links_pend: pend["links"] = links_pend

    # registry ids: verified ones switch on the bio.tools / GitHub data on the live page; the reviewers
    # put citations in the DOI column
    ids, ids_pend = {}, {}
    for name, key in [("bio.tools id", "biotools"), ("GitHub repo", "github"), ("RSD ID", "rsd")]:
        c = cell(name); st = status(c)
        if st and st != "struck":
            (ids if st == "verified" else ids_pend)[key] = str(c.value).strip()
            if st == "decide": report["decide"].append((tid, name, c.value))
    c = cell("DOI"); st = status(c)
    if st and st != "struck":
        v = str(c.value).strip()
        m = re.search(r"(10\.\d{4,9}/[^\s]+)", v)
        if m: (ids if st == "verified" else ids_pend)["doi"] = m.group(1).rstrip(".")
        if len(v) > 60 or not m:
            (top if st == "verified" else pend)["how_to_cite"] = v
        if st == "decide": report["decide"].append((tid, "DOI", v))
    if (prev.get("ids") or {}).get("fairsharing"):
        ids_pend["fairsharing"] = prev["ids"]["fairsharing"]     # column was removed from the sheet
    if ids: top["ids"] = ids
    if ids_pend: pend["ids"] = ids_pend
    multi(cols("Keyword 1", "Keyword 6"), "keywords", lambda v: KEYWORDS.get(str(v).strip().lower(), str(v).strip()))
    # related tools are no longer in the sheet: keep the current ones, verified or not
    for block, cur in ((top, current.get(tid, {})), (pend, current.get(tid, {}).get("pending", {}))):
        if cur.get("related"): block["related"] = list(cur["related"])
    # keep related only for tools that stay in the catalogue (checked after the loop)
    if prev.get("institutes"):
        pend["institutes"] = prev["institutes"]                  # Used by was not part of this review
    for field, value in CHANGES.get(tid, {}).items():
        if field.startswith("links."):
            key, block, pblock = field[6:], top.setdefault("links", {}), pend.get("links", {})
        else:
            key, block, pblock = field, top, pend
        report["overrides"].append((tid, field, block.get(key, pblock.get(key)), value))
        block[key] = value; pblock.pop(key, None)
    if not pend.get("links", True): del pend["links"]
    if pend: top["pending"] = pend
    entries.append(top)

kept = {e["id"] for e in entries}
for e in entries:
    for block in (e, e.get("pending", {})):
        if "related" in block:
            block["related"] = [x for x in block["related"] if x in kept]
            if not block["related"]: del block["related"]

# ---------------------------------------------------------------- YAML output
ORDER = ["id", "name", "redirect_from", "short_description", "description", "tool_type", "fair", "lifecycle", "domains", "license",
         "cost", "code_availability", "maturity", "deprecated", "hosted", "links", "logo", "keywords", "ids",
         "how_to_cite", "institutes", "related", "pending"]
LINK_ORDER = ["homepage", "homepage_opens_tool", "docs", "docs_note", "download", "download_note", "other"]

def q(s):
    s = str(s)
    if s == "" or re.search(r"[:#\[\]{},&*!|>'\"%@`]|^\s|\s$|^(yes|no|true|false|null|on|off)$", s, re.I) or s[0] in "-?":
        return json.dumps(s, ensure_ascii=False)
    return s

def emit(d, indent, order):
    out = []
    pad = " " * indent
    for k in sorted(d, key=lambda x: order.index(x) if x in order else 99):
        v = d[k]
        if k == "pending":
            out.append(f"{pad}pending:   # not verified yet: hidden on the live site, shown with ?unverified=1")
            out += emit(v, indent + 2, order); continue
        if k == "links":
            out.append(f"{pad}links:"); out += emit(v, indent + 2, LINK_ORDER); continue
        if k == "ids":
            out.append(f"{pad}ids:")
            out += [f"{pad}  {ik}: {q(iv)}" for ik, iv in v.items()]; continue
        if isinstance(v, bool):
            out.append(f"{pad}{k}: {'true' if v else 'false'}")
        elif isinstance(v, list) and v and isinstance(v[0], dict):
            out.append(f"{pad}{k}:")
            for item in v:
                inner = ", ".join(f"{ik}: {q(iv)}" for ik, iv in item.items())
                out.append(f"{pad}  - {{{inner}}}")
        elif isinstance(v, list):
            out.append(f"{pad}{k}: [{', '.join(q(x) for x in v)}]")
        elif k in ("description", "how_to_cite") and len(str(v)) > 90:
            out.append(f"{pad}{k}: >-")
            words, line = str(v).split(), ""
            for w in words:
                if len(line) + len(w) + 1 > 100:
                    out.append(f"{pad}  {line}"); line = w
                else:
                    line = f"{line} {w}".strip()
            if line: out.append(f"{pad}  {line}")
        else:
            out.append(f"{pad}{k}: {q(v)}")
    return out

HEADER = f"""# Tool entries (schema: _schemas/tool.schema.yml), rendered with _layouts/tool_profile.html.
#
# Imported on {datetime.date.today()} from the manual review (_review/{XLSX.name}):
# - top-level values were verified by a reviewer (green or yellow in the sheet);
# - `pending:` holds values nobody has checked yet (white) or that still need a decision (purple).
#   The live site shows only verified values and hides empty sections; add ?unverified=1 to a tool
#   page URL to see the pending values too, with a dashed border.
#   When a pending value is checked, move it up to the top level.
# - values and tools that were crossed out in the sheet are not included.
#
# tool_type: bio.tools toolType values, or Self-assessment / checklist, Registry / catalogue,
#            Standard / specification (_data/tool_types.yml).
# fair:      F, A, I, R and/or planning. domains: _data/domains.yml. keywords: _data/keywords.yml.
# code_availability: one value, or two for tools with an open-source and a commercial edition.
# links.homepage_opens_tool: true when the homepage link opens the working tool itself.
# hosted:    true for online services.
"""
text = HEADER + "\n" + "\n\n".join("- " + "\n".join(emit(e, 2, ORDER))[2:] for e in entries) + "\n"
OUT.write_text(text)
yaml.safe_load(open(OUT))   # must parse
json.dump(report, open(Path(__file__).with_name("import_report.json"), "w"), indent=1, ensure_ascii=False)
print(f"{len(entries)} entries written; removed: {report['removed_tools']}; struck values: {len(report['struck_values'])}; "
      f"undecided: {len(report['decide'])}; notes mapped: {len(report['notes_mapped'])}")
